# -*- coding: utf-8 -*-
"""스킬 마법사용 엑셀 도구.

사용법:
	python xlsx_tool.py inspect                      문답에 필요한 enum · 기존 ID · SkillParamDef 출력
	python xlsx_tool.py append <spec.json> [--dry-run]  spec 의 행을 검증한 뒤 각 시트 끝에 추가
	python xlsx_tool.py build                        ExcelDataTable.exe --build 를 이 체크아웃 경로로 실행

공통 옵션:
	--root <경로>   데이터(xlsx · 생성물) 기준 저장소 루트. 기본값은 이 스크립트가 있는 체크아웃.
	                변환기 exe 는 항상 이 스크립트가 있는 체크아웃의 것을 쓴다.

spec.json 형식:
	{"rows": [{"file": "Skill.xlsx", "sheet": "#Skill", "values": {"ID": "Skill_X", "CastingParam": 1.5}}]}
	- 빈 칸은 키를 빼거나 null 로 둔다
	- 숫자는 JSON 숫자로 쓴다 (!string 칸이라도 숫자면 숫자 셀로 기록 — 기존 데이터와 같은 형태)
	- 배열 컬럼(edt_settings.json IsArray)은 리스트 또는 "A;B" 문자열
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SCRIPT_ROOT = Path(__file__).resolve().parents[4]
PRIMITIVE_TYPES = ("!int", "!float", "!bool", "!string")
IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
LOCK_HINT = "해당 xlsx 를 Excel 에서 닫을 것 (Excel 이 꺼져 있는데 남아 있으면 오래된 잔재 — ~$ 파일을 지우면 된다)"

# inspect 가 enum 블록을 출력할 파일, 그 외 파일에서 따로 뽑을 enum
INSPECT_ENUM_FILES = ("Skill.xlsx", "Buff.xlsx", "Summon.xlsx")
INSPECT_EXTRA_ENUMS = {"Equipment.xlsx": ("WeaponType",)}


def getXlsDir(root):
	return root / "Shared" / "XLS"


def findLockFiles(xlsDir):
	return sorted(p.name for p in xlsDir.glob("~$*.xlsx"))


def listXlsxFiles(xlsDir):
	return sorted(p for p in xlsDir.glob("*.xlsx") if not p.name.startswith("~$"))


def cellText(value):
	if value is None:
		return ""
	return str(value).strip()


def readEnumBlocks(ws):
	"""!Enum 시트의 블록들을 [{name, note, members:[{name, desc, note}]}] 로 읽는다."""
	blocks = []
	current = None
	for row in ws.iter_rows(values_only=True):
		cells = list(row) + [None] * 5
		marker = cellText(cells[0])
		name = cellText(cells[1])
		if marker == "!Enum":
			current = {"name": name, "note": cellText(cells[4]), "members": []}
			blocks.append(current)
			continue

		if current is not None and name:
			current["members"].append({"name": name, "desc": cellText(cells[2]), "note": cellText(cells[4])})

	return blocks


def readTableSheet(ws):
	"""!Table / !EnumTable 시트의 헤더 · 기존 ID · 마지막 데이터 행을 읽는다."""
	rows = list(ws.iter_rows(values_only=True))
	info = {"marker": cellText(rows[0][0]) if rows else "", "tableName": cellText(rows[0][1]) if rows and len(rows[0]) > 1 else ""}
	headerIndex = -1
	for i in range(len(rows)):
		if len(rows[i]) > 1 and cellText(rows[i][1]) == "!ID":
			headerIndex = i
			break

	if headerIndex < 0 or headerIndex + 1 >= len(rows):
		info["error"] = "!ID 헤더 행을 찾지 못함"
		return info

	nameRow = rows[headerIndex]
	typeRow = rows[headerIndex + 1]
	# B열 = ID, C열부터 처음 빈 헤더 칸 전까지가 필드 (변환기와 같은 규칙)
	columns = [{"name": "ID", "type": cellText(typeRow[1]), "col": 2}]
	for c in range(2, len(nameRow)):
		fieldName = cellText(nameRow[c])
		if not fieldName:
			break

		columns.append({"name": fieldName, "type": cellText(typeRow[c]) if c < len(typeRow) else "", "col": c + 1})

	dataRows = []
	lastRow = headerIndex + 2
	for i in range(headerIndex + 2, len(rows)):
		idText = cellText(rows[i][1]) if len(rows[i]) > 1 else ""
		if not idText:
			continue

		lastRow = i + 1
		record = {}
		for column in columns:
			index = column["col"] - 1
			record[column["name"]] = rows[i][index] if index < len(rows[i]) else None

		dataRows.append(record)

	info["columns"] = columns
	info["rows"] = dataRows
	info["ids"] = [cellText(r["ID"]) for r in dataRows]
	info["lastRow"] = lastRow
	return info


def collectTypeMembers(xlsDir):
	"""모든 xlsx 의 !Enum 블록과 !EnumTable ID 를 {타입명: set(멤버)} 로 모은다."""
	members = {}
	for path in listXlsxFiles(xlsDir):
		wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
		for ws in wb.worksheets:
			marker = cellText(ws.cell(1, 1).value)
			if marker == "!Enum":
				for block in readEnumBlocks(ws):
					members.setdefault(block["name"], set()).update(m["name"] for m in block["members"])

			elif marker == "!EnumTable":
				info = readTableSheet(ws)
				if "error" not in info:
					members.setdefault(info["tableName"], set()).update(info["ids"])

		wb.close()

	return members


def loadArrayFields(xlsDir):
	settingsPath = xlsDir / "edt_settings.json"
	settings = json.loads(settingsPath.read_text(encoding="utf-8"))
	fields = settings.get("FieldSettings", {})
	return {key for key in fields if fields[key].get("IsArray")}


def formatNumber(value):
	if isinstance(value, float) and value.is_integer():
		return str(int(value))

	return str(value)


# ───────────────────────── inspect ─────────────────────────

def printEnumBlock(block):
	parts = []
	for member in block["members"]:
		extra = "; ".join(x for x in (member["desc"], member["note"]) if x)
		parts.append(f"{member['name']}({extra})" if extra else member["name"])

	header = f"{block['name']} [{block['note']}]" if block["note"] else block["name"]
	print(f"{header}: " + " | ".join(parts))


def runInspect(root):
	xlsDir = getXlsDir(root)
	locks = findLockFiles(xlsDir)
	print("# 잠금 파일(~$): " + (", ".join(locks) if locks else "없음"))

	print("\n# enum")
	for fileName in INSPECT_ENUM_FILES:
		wb = openpyxl.load_workbook(xlsDir / fileName, read_only=True, data_only=True)
		for ws in wb.worksheets:
			if cellText(ws.cell(1, 1).value) == "!Enum":
				for block in readEnumBlocks(ws):
					printEnumBlock(block)

		wb.close()

	for fileName in INSPECT_EXTRA_ENUMS:
		wanted = INSPECT_EXTRA_ENUMS[fileName]
		wb = openpyxl.load_workbook(xlsDir / fileName, read_only=True, data_only=True)
		for ws in wb.worksheets:
			if cellText(ws.cell(1, 1).value) == "!Enum":
				for block in readEnumBlocks(ws):
					if block["name"] in wanted:
						printEnumBlock(block)

		wb.close()

	print("\n# 기존 ID")
	skillWb = openpyxl.load_workbook(xlsDir / "Skill.xlsx", read_only=True, data_only=True)
	skill = readTableSheet(skillWb["#Skill"])
	print(f"Skill ({len(skill['rows'])}): " + ", ".join(
		f"{cellText(r['ID'])}[{cellText(r['SkillCategory']) or '-'}/{cellText(r['CastingType']) or '-'}]" for r in skill["rows"]))
	effect = readTableSheet(skillWb["#SkillEffect"])
	print(f"SkillEffect ({len(effect['rows'])}): " + ", ".join(
		f"{cellText(r['ID'])}[{cellText(r['EffectType'])}]" for r in effect["rows"]))
	paramDef = readTableSheet(skillWb["#SkillParamDef"])
	skillWb.close()

	buffWb = openpyxl.load_workbook(xlsDir / "Buff.xlsx", read_only=True, data_only=True)
	buff = readTableSheet(buffWb["#Buff"])
	buffWb.close()
	print(f"Buff ({len(buff['rows'])}): " + ", ".join(f"{cellText(r['ID'])}({cellText(r['Name'])})" for r in buff["rows"]))

	projectileWb = openpyxl.load_workbook(xlsDir / "Projectile.xlsx", read_only=True, data_only=True)
	projectile = readTableSheet(projectileWb["#Projectile"])
	projectileWb.close()
	print(f"Projectile ({len(projectile['rows'])}): " + ", ".join(projectile["ids"]))

	summonWb = openpyxl.load_workbook(xlsDir / "Summon.xlsx", read_only=True, data_only=True)
	summon = readTableSheet(summonWb["#Summon"])
	summonWb.close()
	print(f"Summon ({len(summon['rows'])}): " + ", ".join(
		f"{cellText(r['ID'])}[{cellText(r['AIType'])}]" for r in summon["rows"]))

	statWb = openpyxl.load_workbook(xlsDir / "Stat.xlsx", read_only=True, data_only=True)
	stat = readTableSheet(statWb["#Stat"])
	statDetail = readTableSheet(statWb["#StatDetail"])
	statWb.close()
	print(f"Stat ({len(stat['rows'])}): " + ", ".join(f"{cellText(r['ID'])}({cellText(r['Name'])})" for r in stat["rows"]))
	print(f"StatDetail ({len(statDetail['rows'])}): " + ", ".join(statDetail["ids"]))

	print("\n# SkillParamDef (EffectType별 EffectParam_1~5)")
	byType = {}
	for r in paramDef["rows"]:
		byType.setdefault(cellText(r["EffectType"]), []).append(r)

	for effectType in byType:
		parts = []
		for r in byType[effectType]:
			key = cellText(r["ParamKey"])
			if not key:
				parts.append(f"{formatNumber(r['Index'])} (미사용)")
				continue

			extra = "; ".join(x for x in (cellText(r["ValueType"]), cellText(r["Desc"])) if x)
			parts.append(f"{formatNumber(r['Index'])} {key}({extra})")

		print(f"{effectType}: " + " | ".join(parts))

	return 0


# ───────────────────────── append ─────────────────────────

def convertValue(rawValue, column, tableName, arrayFields, typeMembers, errors, label):
	"""spec 값을 컬럼 타입에 맞춰 셀 값으로 바꾼다. 빈 칸이면 None."""
	if rawValue is None or (isinstance(rawValue, str) and rawValue.strip() == ""):
		return None

	typeName = column["type"]
	if typeName == "!string":
		if isinstance(rawValue, bool):
			errors.append(f"{label}: !string 칸에 bool 값")
			return None

		if isinstance(rawValue, (int, float)):
			return rawValue

		return str(rawValue)

	if typeName == "!int":
		if isinstance(rawValue, bool):
			errors.append(f"{label}: !int 칸에 bool 값")
			return None

		text = cellText(rawValue)
		if not re.fullmatch(r"-?\d+(\.0+)?", text):
			errors.append(f"{label}: 정수가 아님 ({text})")
			return None

		return int(float(text))

	if typeName == "!float":
		if isinstance(rawValue, bool):
			errors.append(f"{label}: !float 칸에 bool 값")
			return None

		text = cellText(rawValue)
		if not re.fullmatch(r"-?(\d+(\.\d*)?|\.\d+)([eE][-+]?\d+)?", text):
			errors.append(f"{label}: 숫자가 아님 ({text})")
			return None

		number = float(text)
		return int(number) if number.is_integer() else number

	if typeName == "!bool":
		if isinstance(rawValue, bool):
			return rawValue

		text = cellText(rawValue).lower()
		if text in ("true", "1"):
			return True

		if text in ("false", "0"):
			return False

		errors.append(f"{label}: bool 이 아님 ({rawValue})")
		return None

	# 나머지 !Xxx = enum / EnumTable 참조
	refType = typeName[1:]
	if refType not in typeMembers:
		errors.append(f"{label}: 알 수 없는 타입 {typeName}")
		return None

	isArray = f"{tableName}.{column['name']}" in arrayFields
	if isinstance(rawValue, list):
		items = [cellText(v) for v in rawValue]
	elif isArray:
		items = [cellText(v) for v in str(rawValue).split(";")]
	else:
		items = [cellText(rawValue)]

	items = [v for v in items if v and v != "None"]
	if not items:
		return None

	if len(items) > 1 and not isArray:
		errors.append(f"{label}: 배열 컬럼이 아닌데 값이 여러 개 ({items})")
		return None

	for item in items:
		if item not in typeMembers[refType]:
			errors.append(f"{label}: {refType} 에 없는 값 '{item}'")

	return ";".join(items)


def runAppend(root, specPath, dryRun):
	xlsDir = getXlsDir(root)
	spec = json.loads(Path(specPath).read_text(encoding="utf-8"))
	specRows = spec.get("rows", [])
	errors = []
	if not specRows:
		print("오류: spec 에 rows 가 없다")
		return 1

	# 쓰기 대상 파일의 잠금만 막는다 (무관한 파일의 오래된 ~$ 잔재로 멈추지 않게)
	targetLocks = {"~$" + r.get("file", "") for r in specRows}
	locks = [name for name in findLockFiles(xlsDir) if name in targetLocks]
	if locks:
		errors.append(f"잠금 파일 존재: {', '.join(locks)} — {LOCK_HINT}")

	typeMembers = collectTypeMembers(xlsDir)
	arrayFields = loadArrayFields(xlsDir)

	# 1) 대상 시트 정보 로드 + 신규 ID 선등록 (spec 안의 상호 참조 허용)
	sheetInfos = {}
	pendingIds = {}
	validIds = {}
	for index in range(len(specRows)):
		specRow = specRows[index]
		fileName = specRow.get("file", "")
		sheetName = specRow.get("sheet", "")
		values = specRow.get("values", {})
		label = f"rows[{index}] {fileName} {sheetName}"
		key = (fileName, sheetName)
		if key not in sheetInfos:
			path = xlsDir / fileName
			if not path.is_file():
				errors.append(f"{label}: 파일 없음")
				sheetInfos[key] = None
				continue

			wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
			if sheetName not in wb.sheetnames:
				errors.append(f"{label}: 시트 없음 (있는 시트: {', '.join(wb.sheetnames)})")
				sheetInfos[key] = None
				wb.close()
				continue

			info = readTableSheet(wb[sheetName])
			wb.close()
			if "error" in info:
				errors.append(f"{label}: {info['error']}")
				sheetInfos[key] = None
				continue

			if info["marker"] != "!EnumTable":
				errors.append(f"{label}: !EnumTable 시트만 지원 (현재 {info['marker']})")
				sheetInfos[key] = None
				continue

			sheetInfos[key] = info

		info = sheetInfos[key]
		if info is None:
			continue

		newId = cellText(values.get("ID"))
		if not IDENTIFIER_PATTERN.fullmatch(newId):
			errors.append(f"{label}: ID 는 영문·숫자·_ 만 (현재 '{newId}')")
			continue

		tableName = info["tableName"]
		if newId in info["ids"]:
			errors.append(f"{label}: 이미 있는 ID {newId}")
			continue

		if newId in pendingIds.setdefault(tableName, set()):
			errors.append(f"{label}: spec 안에서 ID 중복 {newId}")
			continue

		pendingIds[tableName].add(newId)
		typeMembers.setdefault(tableName, set()).add(newId)
		validIds[index] = newId

	# 2) 값 변환 · 검증
	plans = []
	for index in range(len(specRows)):
		specRow = specRows[index]
		key = (specRow.get("file", ""), specRow.get("sheet", ""))
		info = sheetInfos.get(key)
		if info is None or index not in validIds:
			continue

		values = specRow.get("values", {})
		columnByName = {c["name"]: c for c in info["columns"]}
		cells = {"ID": validIds[index]}
		for name in values:
			if name == "ID":
				continue

			label = f"rows[{index}] {validIds[index]}.{name}"
			if name not in columnByName:
				errors.append(f"{label}: 없는 컬럼 (있는 컬럼: {', '.join(columnByName)})")
				continue

			converted = convertValue(values[name], columnByName[name], info["tableName"], arrayFields, typeMembers, errors, label)
			if converted is not None:
				cells[name] = converted

		plans.append({"key": key, "info": info, "cells": cells})

	if errors:
		print("검증 실패 — 아무 파일도 쓰지 않았다")
		for error in errors:
			print("  - " + error)

		return 1

	# 3) 시트별 기록 행 배정 (마지막 데이터 행 다음부터)
	nextRow = {}
	for plan in plans:
		key = plan["key"]
		nextRow[key] = nextRow.get(key, plan["info"]["lastRow"]) + 1
		plan["row"] = nextRow[key]

	for plan in plans:
		cells = plan["cells"]
		detail = ", ".join(f"{name}={formatNumber(cells[name])}" for name in cells if name != "ID")
		prefix = "[dry-run] " if dryRun else ""
		print(f"{prefix}{plan['key'][0]} {plan['key'][1]} {plan['row']}행 {cells['ID']}: {detail}")

	if dryRun:
		print("[dry-run] 검증 통과 — 파일은 쓰지 않았다")
		return 0

	# 4) 파일별로 한 번만 열어 기록
	fileNames = []
	for plan in plans:
		if plan["key"][0] not in fileNames:
			fileNames.append(plan["key"][0])

	for fileName in fileNames:
		path = xlsDir / fileName
		wb = openpyxl.load_workbook(path)
		for plan in plans:
			if plan["key"][0] != fileName:
				continue

			ws = wb[plan["key"][1]]
			columnByName = {c["name"]: c for c in plan["info"]["columns"]}
			for name in plan["cells"]:
				ws.cell(plan["row"], columnByName[name]["col"], plan["cells"][name])

		wb.save(path)
		wb.close()

	# 5) 다시 읽어 확인
	failed = 0
	for plan in plans:
		wb = openpyxl.load_workbook(xlsDir / plan["key"][0], read_only=True, data_only=True)
		written = cellText(wb[plan["key"][1]].cell(plan["row"], 2).value)
		wb.close()
		if written != plan["cells"]["ID"]:
			print(f"확인 실패: {plan['key'][0]} {plan['key'][1]} {plan['row']}행 = '{written}' (기대 {plan['cells']['ID']})")
			failed += 1

	if failed:
		return 1

	print(f"추가 완료: {len(plans)}행")
	return 0


# ───────────────────────── build ─────────────────────────

def runBuild(root):
	xlsDir = getXlsDir(root)
	locks = findLockFiles(xlsDir)
	if locks:
		print(f"오류: 잠금 파일 존재: {', '.join(locks)} — {LOCK_HINT}")
		return 1

	sourceExe = SCRIPT_ROOT / "Shared" / "Tools" / "ExcelDataTable" / "ExcelDataTable.exe"
	if not sourceExe.is_file():
		print(f"오류: 변환기 없음 {sourceExe}")
		return 1

	# exe 는 자기 폴더의 edt_local.json 만 읽는다 → 임시 폴더 사본에 이 체크아웃 경로를 넣어 실행
	workDir = Path(tempfile.gettempdir()) / "skill-wizard-edt"
	workDir.mkdir(parents=True, exist_ok=True)
	workExe = workDir / sourceExe.name
	sourceStat = sourceExe.stat()
	if not workExe.is_file() or workExe.stat().st_size != sourceStat.st_size or int(workExe.stat().st_mtime) != int(sourceStat.st_mtime):
		shutil.copy2(sourceExe, workExe)

	localConfig = {
		"DataPath": str(xlsDir),
		"DataOutputPath": str(root / "ProjectOne" / "Assets" / "Project" / "Data" / "Tables"),
		"ScriptOutputPath": str(root / "ProjectOne" / "Assets" / "Project" / "Scripts" / "Core" / "ExcelData"),
	}
	(workDir / "edt_local.json").write_text(json.dumps(localConfig, ensure_ascii=False, indent=2), encoding="utf-8")
	print("빌드 경로: " + json.dumps(localConfig, ensure_ascii=False))

	result = subprocess.run([str(workExe), "--build"], cwd=workDir, capture_output=True)
	print(result.stdout.decode("utf-8", errors="replace"))
	if result.stderr:
		print(result.stderr.decode("utf-8", errors="replace"))

	print(f"exit code: {result.returncode}")
	return result.returncode


def main():
	common = argparse.ArgumentParser(add_help=False)
	common.add_argument("--root", default=str(SCRIPT_ROOT), help="데이터 기준 저장소 루트")
	parser = argparse.ArgumentParser(description="스킬 마법사용 엑셀 도구")
	commands = parser.add_subparsers(dest="command", required=True)
	commands.add_parser("inspect", parents=[common])
	appendParser = commands.add_parser("append", parents=[common])
	appendParser.add_argument("spec")
	appendParser.add_argument("--dry-run", action="store_true")
	commands.add_parser("build", parents=[common])
	args = parser.parse_args()

	root = Path(args.root).resolve()
	if args.command == "inspect":
		return runInspect(root)

	if args.command == "append":
		return runAppend(root, args.spec, args.dry_run)

	return runBuild(root)


if __name__ == "__main__":
	sys.exit(main())
